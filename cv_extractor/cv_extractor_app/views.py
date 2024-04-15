import os
import re
import PyPDF2
import pandas as pd
from pandas import ExcelWriter
from django.http import HttpRequest
from django.views.generic import TemplateView
from django.shortcuts import render
from docx import Document
from .models import extractor
from openpyxl import load_workbook
import shutil

class extractorgeneric(TemplateView):
    template_name = 'input.html'

    def post(self, request: HttpRequest, *args, **kwargs):
        try:
            datas = request.FILES.getlist('cv')
            for datas in datas:
                try:
                    upload_data = extractor(uploaded_file=datas)
                    upload_data.save()
                except Exception as e:
                    # Handle specific exceptions related to extractor, if any
                    print(f"Error occurred while processing file: {e}")

            # Function to extract email IDs and contact numbers from text
            def extract_contacts(text):
                emails = re.findall(r'\S+@\S+', text)
                phone_numbers = re.findall(r'(\d{3}[-\.\s]??\d{3}[-\.\s]??\d{4}|\(\d{3}\)\s*\d{3}[-\.\s]??\d{4}|\d{3}[-\.\s]??\d{4})', text)
                return emails, phone_numbers

            # Function to extract text from PDF
            def extract_text_from_pdf(pdf_path):
                text = ''
                with open(pdf_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page_num in range(len(pdf_reader.pages)):
                        page = pdf_reader.pages[page_num]
                        text += page.extract_text()
                return text

            # Function to extract text from DOC file
            def extract_text_from_docx(doc_path):
                doc = Document(doc_path)
                text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
                return text
            
            def delete_files_and_instances():
                print("entered")
                # Path to the directory containing uploaded files
                upload_path = r'D:\Django\resume extractor\cv_extractor\resumes'

                # Query all instances of your model
                instances = extractor.objects.all()

                # Iterate over the instances
                for instance in instances:
                    instance.delete()
                shutil.rmtree(upload_path)

            # Path to the directory containing CVs
            cv_directory = r'D:\Django\resume extractor\cv_extractor\resumes'

            # List to store extracted data
            data = []

            # Iterate over each file in the directory
            for file_name in os.listdir(cv_directory):
                file_path = os.path.join(cv_directory, file_name)
                try:
                    if file_name.endswith('.pdf'):
                        text = extract_text_from_pdf(file_path)
                    elif file_name.endswith('.docx'):
                        text = extract_text_from_docx(file_path)
                    else:
                        continue
                except Exception as e:
                    # Handle specific exceptions related to text extraction, if any
                    msg = "Error occurred while extracting text from file"
                    continue

                emails, phone_numbers = extract_contacts(text)
                data.append({'File': file_name, 'Email': emails, 'Phone Number': phone_numbers, 'Text': text})

            output_path = r"D:\Django\resume extractor\cv_extractor\static\output_excel\parsed_data.xlsx"

            # Create a DataFrame
            df = pd.DataFrame(data)

            try:
                df.to_excel(output_path, index=False)  # Write new data to the Excel file
                msg = "Successfully parsed entered CV's"
                delete_files_and_instances()
            except Exception as e:
                # Handle specific exceptions related to Excel saving, if any
                msg = "Error occurred while saving Excel file"
                print(e)


            return render(request, 'input.html', {'msg':msg})
        except Exception as e:
            # Handle any unexpected exceptions
            msg ="Error"
            return render(request, 'input.html', {'msg':msg})
